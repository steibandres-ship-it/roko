"""Bounded Excel import with preview and transactional, conflict-safe insertion."""
from io import BytesIO
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile, BadZipFile
from urllib.parse import unquote

from email_validator import validate_email, EmailNotValidError
from openpyxl import load_workbook
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import FileResponse
from starlette.concurrency import run_in_threadpool
from sqlalchemy import select, func
from sqlalchemy.orm import Session

from .db import get_session
from .fans import local_operator
from .preuser_models import Preuser, PreuserImport

MAX_ROWS = 50_000
MAX_BYTES = 10 * 1024 * 1024
router = APIRouter()
admin = APIRouter(prefix='/api/preusers', dependencies=[Depends(local_operator)], tags=['preusers'])


def parse_excel(data: bytes):
    if len(data) > MAX_BYTES:
        raise HTTPException(413, 'El archivo supera 10 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 64 * 1024 * 1024 or len(archive.infolist()) > 1000:
                raise HTTPException(413, 'El contenido descomprimido supera el límite permitido.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'No se pudo leer el Excel. Usa un archivo .xlsx válido, sin contraseña.') from None
    emails, errors, seen = [], [], set()
    invalid = duplicates = rows = blank = 0
    try:
        if len(book.worksheets) != 1:
            raise HTTPException(422, 'Usa un Excel de una sola hoja y una sola columna de correos.')
        sheet = book.worksheets[0]
        if (sheet.max_column or 0) > 10 or (sheet.max_row or 0) > MAX_ROWS + 1:
            raise HTTPException(422, 'Máximo 50.000 filas y una columna; elimina filas o columnas extra con formato.')
        # Do not trust worksheet dimension metadata to delimit the import.
        sheet.reset_dimensions()
        for number, cells in enumerate(sheet.iter_rows(), 1):
            if number > MAX_ROWS + 1:
                raise HTTPException(422, 'Máximo 50.000 correos por archivo, más un encabezado opcional.')
            if any(c.value is not None and str(c.value).strip() for c in cells[1:]):
                raise HTTPException(422, f'Fila {number}: el archivo debe contener solamente la columna de correo.')
            cell = cells[0] if cells else None
            value = cell.value if cell else None
            if number == 1 and isinstance(value, str) and value.strip().casefold() in ('email', 'correo', 'correo electrónico', 'correo electronico'):
                continue
            if value is None or (isinstance(value, str) and not value.strip()):
                blank += 1
                continue
            rows += 1
            if rows > MAX_ROWS:
                raise HTTPException(422, 'Máximo 50.000 correos por archivo.')
            try:
                if cell.data_type == 'f' or not isinstance(value, str):
                    raise ValueError()
                email = validate_email(value.strip(), check_deliverability=False, allow_smtputf8=False).normalized.lower()
                if len(email) > 254:
                    raise ValueError()
            except (ValueError, EmailNotValidError):
                invalid += 1
                if len(errors) < 100:
                    errors.append({'row': number, 'reason': 'Correo inválido, fórmula o valor no textual.'})
                continue
            if email in seen:
                duplicates += 1
            else:
                seen.add(email)
                emails.append(email)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'El archivo tiene contenido no válido. Exporta nuevamente como .xlsx.') from None
    finally:
        book.close()
    if not rows:
        raise HTTPException(422, 'El archivo no contiene correos.')
    return emails, {'rows': rows, 'invalid': invalid, 'duplicates_in_file': duplicates, 'blank_rows': blank, 'errors': errors}


def import_excel(data, filename, preview, session):
    emails, report = parse_excel(data)
    existing = 0
    for offset in range(0, len(emails), 500):
        existing += session.scalar(select(func.count()).select_from(Preuser).where(Preuser.email.in_(emails[offset:offset+500])))
    report.update(existing=existing, new=len(emails)-existing, preview=preview)
    if preview:
        return report
    if not emails:
        raise HTTPException(422, 'No hay correos válidos para importar.')
    dialect = session.get_bind().dialect.name
    if dialect == 'postgresql':
        from sqlalchemy.dialects.postgresql import insert
    elif dialect == 'sqlite':
        from sqlalchemy.dialects.sqlite import insert
    else:
        raise HTTPException(503, 'La importación requiere PostgreSQL o SQLite.')
    job = PreuserImport(filename=filename, file_hash=sha256(data).hexdigest(), rows=report['rows'], imported=0, duplicates=0, invalid=report['invalid'], errors=report['errors'])
    try:
        session.add(job)
        session.flush()
        total = 0
        for offset in range(0, len(emails), 500):
            statement = insert(Preuser).values([{'email': email, 'import_id': job.id} for email in emails[offset:offset+500]]).on_conflict_do_nothing(index_elements=['email']).returning(Preuser.id)
            total += len(session.scalars(statement).all())
        job.imported = total
        job.duplicates = report['duplicates_in_file'] + len(emails) - total
        session.commit()
    except Exception:
        session.rollback()
        raise
    report.update(imported=total, existing=len(emails)-total, new=total, import_id=job.id)
    return report


@router.get('/preusers', include_in_schema=False)
def page():
    return FileResponse(Path(__file__).parent / 'static' / 'preusers.html')


@admin.get('')
def overview(session: Session = Depends(get_session)):
    jobs = session.scalars(select(PreuserImport).order_by(PreuserImport.created_at.desc()).limit(20))
    return {'total': session.scalar(select(func.count()).select_from(Preuser)), 'imports': [{'id': j.id, 'filename': j.filename, 'created_at': j.created_at.isoformat(), 'imported': j.imported, 'duplicates': j.duplicates, 'invalid': j.invalid} for j in jobs]}


@admin.post('/import')
async def upload(request: Request, preview: bool = True, session: Session = Depends(get_session)):
    filename = unquote(request.headers.get('x-filename', 'correos.xlsx')).replace('\\', '/').split('/')[-1][:255]
    if not filename.lower().endswith('.xlsx'):
        raise HTTPException(422, 'Sube un archivo Excel .xlsx.')
    body = bytearray()
    async for chunk in request.stream():
        body.extend(chunk)
        if len(body) > MAX_BYTES:
            raise HTTPException(413, 'El archivo supera 10 MB.')
    return await run_in_threadpool(import_excel, bytes(body), filename, preview, session)


router.include_router(admin)
